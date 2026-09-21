"""IDE 执行 / 历史 / 导出路由（I10 设计文档 §4.2，G21 异步执行架构；I4 F16~F18 语义保留）。

- POST /ide/execute：异步任务化（taskId + ThreadPoolExecutor 5 并发，超限 429）
- GET /ide/stream/{taskId}：SSE 推流（log/result/done + 3s 心跳；EventSource ?token= 兜底鉴权）
- POST /ide/cancel/{taskId}：置取消标记 + 独立短连接 KILL QUERY（G21）
- 执行线程：单连接逐条（失败即停、前序保留）；变量渲染（全局参数 by env + 内置时间参数）
- 事务模式 auto=逐语句自动提交 / manual=BEGIN 起步全成功 COMMIT 任一失败 ROLLBACK（G22）
- 历史落库含 log_text（逐条日志 JSON）/ affected_total；sql_text 全文（MEDIUMTEXT）
- POST /ide/lint：sqlglot mysql 方言逐语句语法预检（G13）
- GET /ide/result/{historyId}/{stmtIndex}：按语句重放分页（页间可能漂移，一期接受）
- GET/DELETE /ide/history/{id}：详情（log_text 解析）/ 删除（G5）
- GET /ide/export：按 historyId 重放 SQL 流式导出 CSV（utf8-sig BOM；上限 10 万行）
- GET/POST/PUT/DELETE /ide/scripts：命名脚本 CRUD（用户隔离 + 用户内名称唯一，G16）
"""

import csv
import io
import json
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from decimal import Decimal
from queue import Empty, Queue
from typing import Iterator, Optional

import sqlglot
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from api.auth import ApiError, get_current_user, require_perm
from common.db import get_db, new_session
from common.dsconn import open_connection
from common.log import get_logger
from common.models import DataSource, GlobalParam, IdeHistory, IdeScript, User
from common.resp import (
    DS_NOT_FOUND,
    HISTORY_NOT_FOUND,
    PARAM_INVALID,
    SCRIPT_NOT_FOUND,
    SYSTEM_ERROR,
    TASK_NOT_FOUND,
    PageQuery,
    fmt_dt,
    ok,
    page_result,
)
from common.vars_render import render_text

logger = get_logger("api.ide")

router = APIRouter(prefix="/ide", tags=["ide"])

READ_TIMEOUT = 60  # 执行单语句超时（秒）
FETCH_BATCH = 500  # 批量取行批大小（避免大结果集一次性载入内存）
SCREEN_ROW_CAP = 2000  # 首屏驻留行上限（09-22 裁定：行上限 UI 移除，驻留层仅此不可见常量，rowsTotal 全量计数）
EXPORT_ROW_CAP = 100_000  # 导出上限（§14：超出提示走 C22/同步通道）
BOM = b"\xef\xbb\xbf"

# ---------- 异步执行架构（I10 §4.2 G21/G22） ----------

EXEC_MAX_WORKERS = 5  # 并发执行任务上限（超出 429）
TASK_TTL_SEC = 600  # 完成任务保留时长（结果查看/迟到订阅），过期清理防泄漏
HEARTBEAT_SEC = 3  # SSE 心跳间隔（注释行，防代理断连）
WARNINGS_CAP = 20  # 单语句 SHOW WARNINGS 收集上限
ENV_LIST = ("dev", "staging", "prod")

EXEC_POOL = ThreadPoolExecutor(max_workers=EXEC_MAX_WORKERS, thread_name_prefix="ide-exec")
IDE_TASKS: dict = {}
_TASKS_LOCK = threading.Lock()
_ACTIVE = 0  # 运行中任务数（提交时同步判定，防线程池排队绕过并发上限）


class _Task:
    """内存执行任务记录（状态/事件队列/取消标记/KILL 线程号）。"""

    __slots__ = ("task_id", "ds_id", "db_name", "status", "events", "cancel", "thread_id", "created")

    def __init__(self, task_id: str, ds_id: int, db_name: Optional[str]):
        self.task_id = task_id
        self.ds_id = ds_id
        self.db_name = db_name
        self.status = "running"
        self.events = Queue()
        self.cancel = threading.Event()
        self.thread_id = None  # CONNECTION_ID()，KILL QUERY 用
        self.created = time.monotonic()


def _purge_tasks() -> None:
    """清理 TTL 过期的完成任务（注册表与事件队列及时释放，防内存泄漏）。"""
    now = time.monotonic()
    with _TASKS_LOCK:
        for tid in [
            tid
            for tid, t in IDE_TASKS.items()
            if t.status != "running" and now - t.created > TASK_TTL_SEC
        ]:
            IDE_TASKS.pop(tid, None)


class ExecuteBody(BaseModel):
    datasourceId: int
    db: Optional[str] = None
    sql: str
    mode: str = Field(default="auto")  # auto=逐语句自动提交 / manual=事务包裹（G22）
    env: str = Field(default="dev")  # 全局参数环境分组（G3/G4）


def _split_sql(text: str) -> list:
    """多语句分号拆分（空语句剔除，保序；与 C11 同款口径）。"""
    return [part.strip() for part in (text or "").split(";") if part.strip()]


def _jsonable(value):
    """行值 CSV/JSON 安全化（bytes/datetime/Decimal → str）。"""
    if value is None:
        return ""
    if isinstance(value, (bytes, bytearray)):
        return bytes(value).decode("utf-8", errors="replace")
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return value


def _run_one(cur, stmt: str) -> dict:
    """单语句执行：查询全量计数 + 截断首屏驻留 SCREEN_ROW_CAP；非查询返回影响行数。"""
    cur.execute(stmt)
    if cur.description is None:
        affected = cur.rowcount if cur.rowcount and cur.rowcount > 0 else 0
        return {"kind": "exec", "affected": affected, "columns": [], "rows": [], "rowsTotal": 0}
    columns = [col[0] for col in cur.description]
    rows, total = [], 0
    while True:
        batch = cur.fetchmany(FETCH_BATCH)
        if not batch:
            break
        total += len(batch)
        remain = SCREEN_ROW_CAP - len(rows)
        if remain > 0:
            for row in batch[:remain]:
                rows.append([_jsonable(v) for v in row])
    return {"kind": "query", "affected": 0, "columns": columns, "rows": rows, "rowsTotal": total}


def _fetch_warnings(cur) -> list:
    """SHOW WARNINGS 收集（'Level Code: Message'，上限 WARNINGS_CAP；失败不影响主链）。"""
    try:
        cur.execute("SHOW WARNINGS")
        rows = cur.fetchall()
    except Exception:  # noqa: BLE001
        return []
    return [f"{r[0]} {r[1]}: {str(r[2] or '')}" for r in rows[:WARNINGS_CAP]]


def _exec_worker(task: _Task, payload: dict) -> None:
    """执行线程：单连接逐条（失败即停、前序保留）；渲染→执行→SHOW WARNINGS→首屏入队→历史落库。

    - auto：连接 autocommit(True) 逐语句提交；manual：BEGIN 起步，全成功 COMMIT / 任一失败 ROLLBACK
    - 变量快照明文入日志（09-18 裁定不脱敏）；取消经 cancel 标记 + KILL QUERY 中断
    """
    global _ACTIVE
    events = task.events
    statements = payload["statements"]
    manual = payload["mode"] == "manual"
    entries: list = []  # log_text 逐条
    status = "success"
    error = None
    affected_total = 0
    rows_total_last = 0
    conn = None
    started = time.monotonic()
    sess = new_session()
    try:
        ds = sess.get(DataSource, payload["ds_id"])
        if ds is None:
            raise RuntimeError("数据源不存在或已删除")
        env_params = {
            g.name: (g.value or "")
            for g in sess.query(GlobalParam).filter(GlobalParam.env == payload["env"]).all()
        }
        conn = open_connection(ds, db=payload["db"], read_timeout=READ_TIMEOUT)
        with conn.cursor() as cur:
            cur.execute("SELECT CONNECTION_ID()")
            task.thread_id = int(cur.fetchone()[0])
        if manual:
            conn.begin()
        else:
            conn.autocommit(True)
        for idx, stmt in enumerate(statements):
            if task.cancel.is_set():
                status = "canceled"
                break
            entry = {
                "stmtIndex": idx,
                "sql": stmt,
                "rendered": stmt,
                "status": "ok",
                "kind": "",
                "elapsedMs": 0,
                "warnings": [],
            }
            t0 = time.monotonic()
            try:
                rendered, snapshot = render_text(stmt, env_params)
                entry["rendered"] = rendered
                if snapshot:
                    entry["varSnapshot"] = snapshot
                events.put(
                    {
                        "event": "log",
                        "data": {
                            "stmtIndex": idx,
                            "type": "start",
                            "sql": stmt,
                            "rendered": rendered if rendered != stmt else None,
                            "varSnapshot": snapshot or None,
                        },
                    }
                )
                cur = conn.cursor()
                try:
                    out = _run_one(cur, rendered)
                    out["warnings"] = _fetch_warnings(cur)
                finally:
                    cur.close()
                entry["elapsedMs"] = int((time.monotonic() - t0) * 1000)
                entry["warnings"] = out["warnings"]
                entry["kind"] = out["kind"]
                if out["kind"] == "exec":
                    affected_total += out["affected"]
                    entry["affected"] = out["affected"]
                else:
                    rows_total_last = out["rowsTotal"]
                    entry["rowsTotal"] = out["rowsTotal"]
                events.put(
                    {
                        "event": "log",
                        "data": {
                            "stmtIndex": idx,
                            "type": "end",
                            "kind": out["kind"],
                            "affected": out.get("affected", 0),
                            "rowsTotal": out.get("rowsTotal", 0),
                            "elapsedMs": entry["elapsedMs"],
                            "warnings": out["warnings"],
                        },
                    }
                )
                if out["kind"] == "query":
                    events.put(
                        {
                            "event": "result",
                            "data": {
                                "stmtIndex": idx,
                                "kind": "query",
                                "columns": out["columns"],
                                "rows": out["rows"],
                                "rowsTotal": out["rowsTotal"],
                                "elapsedMs": entry["elapsedMs"],
                            },
                        }
                    )
            except Exception as exc:  # noqa: BLE001 失败即停（含 KILL 中断）
                canceled = task.cancel.is_set()
                entry["status"] = "canceled" if canceled else "error"
                entry["error"] = str(exc)[:2000]
                entry["elapsedMs"] = int((time.monotonic() - t0) * 1000)
                entries.append(entry)
                events.put(
                    {
                        "event": "log",
                        "data": {
                            "stmtIndex": idx,
                            "type": "error",
                            "error": entry["error"],
                            "canceled": canceled,
                            "elapsedMs": entry["elapsedMs"],
                        },
                    }
                )
                if canceled:
                    status = "canceled"
                else:
                    status = "failure"
                    error = entry["error"]
                break
            entries.append(entry)
        if manual and conn is not None:
            try:
                if status == "success":
                    conn.commit()
                    events.put({"event": "log", "data": {"type": "txn", "action": "COMMIT"}})
                else:
                    conn.rollback()
                    events.put(
                        {"event": "log", "data": {"type": "txn", "action": "ROLLBACK", "reason": status}}
                    )
            except Exception as exc:  # noqa: BLE001
                logger.warning("事务收口失败: task=%s %r", task.task_id, exc)
    except Exception as exc:  # noqa: BLE001 连接级失败（数据源缺失/连不上/超时）
        status = "canceled" if task.cancel.is_set() else "failure"
        error = str(exc)[:2000]
        entries.append(
            {
                "stmtIndex": -1,
                "sql": "",
                "rendered": "",
                "status": "error",
                "kind": "",
                "elapsedMs": int((time.monotonic() - started) * 1000),
                "warnings": [],
                "error": error,
            }
        )
        events.put({"event": "log", "data": {"type": "error", "error": error}})
        if manual and conn is not None:
            try:
                conn.rollback()
                events.put({"event": "log", "data": {"type": "txn", "action": "ROLLBACK"}})
            except Exception:  # noqa: BLE001
                pass
    finally:
        if conn is not None:
            try:
                conn.close()
            except Exception:  # noqa: BLE001
                pass
        elapsed = int((time.monotonic() - started) * 1000)
        history_id = None
        # rendered 全文 = 各语句渲染后 SQL 以 `;\n` 拼接（与 _split_sql 分号口径一致，可重拆分重放）
        rendered_sql = ";\n".join(
            e["rendered"]
            for e in entries
            if e.get("stmtIndex", -1) >= 0 and e.get("rendered")
        ) or None
        try:
            hist = IdeHistory(
                user_id=payload["user_id"],
                datasource_id=payload["ds_id"],
                db_name=payload["db"] or None,
                sql_text=payload["sql"],
                rendered_sql=rendered_sql,
                status=status,
                elapsed_ms=elapsed,
                rows_total=rows_total_last,
                affected_total=affected_total,
                error=error,
                log_text=json.dumps(entries, ensure_ascii=False),
            )
            sess.add(hist)
            sess.commit()
            history_id = hist.id
        except Exception as exc:  # noqa: BLE001 历史落库失败不丢执行结果
            sess.rollback()
            logger.error("IDE 历史落库失败: task=%s %r", task.task_id, exc)
        finally:
            sess.close()
        task.status = status
        events.put(
            {
                "event": "done",
                "data": {
                    "status": status,
                    "elapsedMs": elapsed,
                    "historyId": history_id,
                    "affectedTotal": affected_total,
                    "rowsTotal": rows_total_last,
                    "stmtCount": len(statements),
                },
            }
        )
        with _TASKS_LOCK:
            _ACTIVE -= 1
        logger.info(
            "IDE 执行完成: task=%s ds=%s db=%s 语句数=%d → %s %dms",
            task.task_id, payload["ds_id"], payload["db"], len(statements), status, elapsed,
        )


@router.post("/execute")
def ide_execute(
    body: ExecuteBody,
    user: User = Depends(require_perm("run_instance")),
    db: Session = Depends(get_db),
):
    """异步执行：创建任务立即返回 {taskId}；结果经 GET /ide/stream/{taskId} SSE 推送。"""
    global _ACTIVE
    ds = db.get(DataSource, body.datasourceId)
    if ds is None:
        raise ApiError(DS_NOT_FOUND, status=404)
    if ds.type not in ("mysql", "greatdb"):
        raise ApiError(PARAM_INVALID, "IDE 仅支持连接型数据源（mysql/greatdb）", status=400)
    statements = _split_sql(body.sql)
    if not statements:
        raise ApiError(PARAM_INVALID, "SQL 为空", status=400)
    if body.mode not in ("auto", "manual"):
        raise ApiError(PARAM_INVALID, "mode 仅支持 auto/manual", status=400)
    if body.env not in ENV_LIST:
        raise ApiError(PARAM_INVALID, f"env 仅支持 {'/'.join(ENV_LIST)}", status=400)
    _purge_tasks()
    task = _Task(uuid.uuid4().hex, ds.id, body.db or None)
    payload = {
        "user_id": user.id,
        "ds_id": ds.id,
        "db": body.db or None,
        "sql": body.sql,
        "statements": statements,
        "mode": body.mode,
        "env": body.env,
    }
    with _TASKS_LOCK:
        if _ACTIVE >= EXEC_MAX_WORKERS:
            raise ApiError(
                SYSTEM_ERROR,
                f"执行并发已达上限 {EXEC_MAX_WORKERS}，请等待进行中任务完成或先停止",
                status=429,
            )
        _ACTIVE += 1
        IDE_TASKS[task.task_id] = task
    try:
        EXEC_POOL.submit(_exec_worker, task, payload)
    except Exception:
        with _TASKS_LOCK:
            _ACTIVE -= 1
            IDE_TASKS.pop(task.task_id, None)
        raise
    logger.info(
        "IDE 任务提交: task=%s ds=%s db=%s 语句数=%d mode=%s env=%s（操作人 %s）",
        task.task_id, ds.name, body.db, len(statements), body.mode, body.env, user.user_name,
    )
    return ok({"taskId": task.task_id})


@router.get("/stream/{task_id}")
def ide_stream(task_id: str, user: User = Depends(get_current_user)):
    """SSE 推流：log/result/done + 3s 心跳；done 事件后关闭（原生 EventSource 断线自动重连）。"""

    task = IDE_TASKS.get(task_id)
    if task is None:
        raise ApiError(TASK_NOT_FOUND, status=404)

    def gen():
        idle_final = 0
        while True:
            try:
                evt = task.events.get(timeout=HEARTBEAT_SEC)
            except Empty:
                if task.status != "running":
                    idle_final += 1
                    if idle_final >= 2:  # 终态后 done 丢失兜底（理论不发生）
                        break
                yield ": heartbeat\n\n"
                continue
            idle_final = 0
            data = json.dumps(evt["data"], ensure_ascii=False, default=str)
            yield f"event: {evt['event']}\ndata: {data}\n\n"
            if evt["event"] == "done":
                break

    return StreamingResponse(
        gen(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@router.post("/cancel/{task_id}")
def ide_cancel(
    task_id: str,
    user: User = Depends(require_perm("run_instance")),
    db: Session = Depends(get_db),
):
    """停止任务（G21）：置取消标记 + 独立短连接 KILL QUERY；执行线程捕获后落库 canceled。"""
    task = IDE_TASKS.get(task_id)
    if task is None:
        raise ApiError(TASK_NOT_FOUND, status=404)
    if task.status != "running":
        return ok({"canceled": False, "status": task.status})
    task.cancel.set()
    killed = False
    if task.thread_id:
        ds = db.get(DataSource, task.ds_id)
        if ds is not None:
            try:
                conn = open_connection(ds, read_timeout=5)
                try:
                    with conn.cursor() as cur:
                        cur.execute(f"KILL QUERY {int(task.thread_id)}")
                        cur.fetchall()
                    killed = True
                finally:
                    conn.close()
            except Exception as exc:  # noqa: BLE001 KILL 失败不阻断取消（查询可能已结束）
                logger.warning("KILL QUERY 失败: task=%s thread=%s %r", task_id, task.thread_id, exc)
    logger.info("IDE 任务停止: task=%s killed=%s（操作人 %s）", task_id, killed, user.user_name)
    return ok({"canceled": True, "killed": killed})


# ---------- 语法预检（G13） / 结果重放分页（G18） / 历史详情删除（G5） ----------


class LintBody(BaseModel):
    sql: str


def _lint_position(exc: Exception) -> tuple:
    """sqlglot ParseError → (line, col)（无位置信息返回 (None, None)）。"""
    errors = getattr(exc, "errors", None) or []
    if errors and isinstance(errors[0], dict):
        return errors[0].get("line"), errors[0].get("col")
    return None, None


@router.post("/lint")
def ide_lint(body: LintBody, user: User = Depends(get_current_user)):
    """语法预检（G13）：sqlglot mysql 方言逐语句 {stmtIndex, ok, error, line, col}。"""
    out = []
    for idx, stmt in enumerate(_split_sql(body.sql)):
        try:
            sqlglot.parse(stmt, read="mysql")
            out.append({"stmtIndex": idx, "ok": True, "error": None, "line": None, "col": None})
        except Exception as exc:  # noqa: BLE001 解析失败即告警项
            line, col = _lint_position(exc)
            out.append({"stmtIndex": idx, "ok": False, "error": str(exc)[:300], "line": line, "col": col})
    return ok(out)


@router.get("/result/{history_id}/{stmt_index}")
def ide_result(
    history_id: int,
    stmt_index: int,
    offset: int = 0,
    limit: int = 200,
    user: User = Depends(require_perm("run_instance")),
    db: Session = Depends(get_db),
):
    """按语句序号重放分页（G18）：offset/limit 取窗口；非确定性查询页间可能漂移（一期接受）。
    执行 SQL 源 = hist.rendered_sql（渲染后，09-22 加列）或回退原始 sql_text（存量历史）。"""
    if offset < 0 or limit < 1 or limit > SCREEN_ROW_CAP:
        raise ApiError(PARAM_INVALID, f"需满足 offset≥0 且 1≤limit≤{SCREEN_ROW_CAP}", status=400)
    hist = db.get(IdeHistory, history_id)
    if hist is None:
        raise ApiError(HISTORY_NOT_FOUND, status=404)
    ds = db.get(DataSource, hist.datasource_id)
    if ds is None:
        raise ApiError(DS_NOT_FOUND, "数据源已删除，无法重放", status=404)
    statements = _split_sql(hist.rendered_sql or hist.sql_text or "")
    if stmt_index < 0 or stmt_index >= len(statements):
        raise ApiError(PARAM_INVALID, f"语句序号越界（共 {len(statements)} 条）", status=400)
    total_hint = None
    if hist.log_text:  # 新历史优先取落库行数 hint（免全量排干重计，计算优化）
        try:
            for entry in json.loads(hist.log_text):
                if (
                    isinstance(entry, dict)
                    and entry.get("stmtIndex") == stmt_index
                    and entry.get("rowsTotal") is not None
                ):
                    total_hint = int(entry["rowsTotal"])
                    break
        except Exception:  # noqa: BLE001 存量脏数据容错
            total_hint = None
    try:
        conn = open_connection(ds, db=hist.db_name or None, read_timeout=None)  # 重放不限时（对齐导出）
    except Exception as exc:  # noqa: BLE001
        raise ApiError(PARAM_INVALID, f"实例连接失败: {str(exc)[:300]}", status=400) from exc
    try:
        cur = conn.cursor()
        try:
            cur.execute(statements[stmt_index])
            if cur.description is None:
                raise ApiError(PARAM_INVALID, "该语句非查询语句，无结果集可分页", status=400)
            columns = [col[0] for col in cur.description]
            rows, fetched = [], 0
            while len(rows) < limit:
                batch = cur.fetchmany(FETCH_BATCH)
                if not batch:
                    break
                for row in batch:
                    fetched += 1
                    if fetched > offset:
                        rows.append([_jsonable(v) for v in row])
                        if len(rows) >= limit:
                            break
            if total_hint is None:  # I4 存量历史无 log_text：全量排干计数兜底
                while True:
                    batch = cur.fetchmany(FETCH_BATCH)
                    if not batch:
                        break
                    fetched += len(batch)
                total_hint = fetched
        finally:
            cur.close()
    finally:
        conn.close()
    return ok(
        {
            "stmtIndex": stmt_index,
            "columns": columns,
            "rows": rows,
            "offset": offset,
            "limit": limit,
            "rowsTotal": total_hint,
        }
    )


def _hist_payload(row: IdeHistory) -> dict:
    return {
        "id": row.id,
        "datasourceId": row.datasource_id,
        "dbName": row.db_name,
        "sqlText": row.sql_text,
        "status": row.status,
        "elapsedMs": row.elapsed_ms,
        "rowsTotal": row.rows_total,
        "affectedTotal": row.affected_total or 0,
        "exported": bool(row.exported),
        "error": row.error,
        "createTime": fmt_dt(row.create_time),
    }


@router.get("/history")
def ide_history(
    datasourceId: Optional[int] = None,
    keyword: Optional[str] = None,
    page: PageQuery = Depends(),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """执行历史分页（按时间倒序；keyword 匹配 SQL 全文）。"""
    query = db.query(IdeHistory)
    if datasourceId:
        query = query.filter(IdeHistory.datasource_id == datasourceId)
    if keyword:
        query = query.filter(IdeHistory.sql_text.like(f"%{keyword}%"))
    total = query.count()
    rows = query.order_by(IdeHistory.id.desc()).offset(page.offset).limit(page.page_size).all()
    return ok(page_result(total, [_hist_payload(row) for row in rows]))


@router.get("/history/{history_id}")
def ide_history_detail(
    history_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """历史详情（G5 查看）：全文 SQL + log_text 解析（逐条语句/耗时/警告/变量快照）。"""
    row = db.get(IdeHistory, history_id)
    if row is None:
        raise ApiError(HISTORY_NOT_FOUND, status=404)
    payload = _hist_payload(row)
    log_entries = []
    if row.log_text:
        try:
            parsed = json.loads(row.log_text)
            if isinstance(parsed, list):
                log_entries = parsed
        except Exception:  # noqa: BLE001 存量脏数据容错
            log_entries = []
    payload["logText"] = log_entries
    return ok(payload)


@router.delete("/history/{history_id}")
def ide_history_delete(
    history_id: int,
    user: User = Depends(require_perm("run_instance")),
    db: Session = Depends(get_db),
):
    """删除执行历史（G5；二次确认由前端承担）。"""
    row = db.get(IdeHistory, history_id)
    if row is None:
        raise ApiError(HISTORY_NOT_FOUND, status=404)
    db.delete(row)
    db.commit()
    logger.info("IDE 历史删除: id=%d（操作人 %s）", history_id, user.user_name)
    return ok(True)


def _csv_stream(ds: DataSource, db_name: Optional[str], sql_text: str) -> Iterator[bytes]:
    """重放 SQL 流式导出 CSV：首个查询结果前加 BOM（防 Excel 乱码）；
    多查询语句各成段落（空行分隔）；全局 10 万行上限。"""
    conn = open_connection(ds, db=db_name or None, read_timeout=None)  # 导出不受执行超时限制
    total_rows = 0
    first_chunk = True
    try:
        cur = conn.cursor()
        try:
            first_section = True
            for stmt in _split_sql(sql_text):
                if total_rows >= EXPORT_ROW_CAP:
                    break
                cur.execute(stmt)
                if cur.description is None:
                    continue  # 非查询语句跳过
                buf = io.StringIO()
                writer = csv.writer(buf)
                if not first_section:
                    writer.writerow([])  # 段间空行
                writer.writerow([col[0] for col in cur.description])
                first_section = False
                while total_rows < EXPORT_ROW_CAP:
                    batch = cur.fetchmany(FETCH_BATCH)
                    if not batch:
                        break
                    for row in batch:
                        writer.writerow([_jsonable(v) for v in row])
                        total_rows += 1
                        if total_rows >= EXPORT_ROW_CAP:
                            break
                    chunk = buf.getvalue().encode("utf-8")
                    yield BOM + chunk if first_chunk else chunk  # BOM 仅首个分片
                    first_chunk = False
                    buf = io.StringIO()
                    writer = csv.writer(buf)
                yield buf.getvalue().encode("utf-8")  # flush 剩余（header/尾批）
        finally:
            cur.close()
        if total_rows >= EXPORT_ROW_CAP:
            yield f"# 超出导出上限 {EXPORT_ROW_CAP} 行，请改用 C22 文件读取/同步通道".encode("utf-8")
    finally:
        conn.close()


def _json_stream(ds: DataSource, db_name: Optional[str], sql_text: str) -> Iterator[bytes]:
    """重放 SQL 流式导出 JSON（G20）：{"sections":[{columns,rows},...],"truncated":bool}。
    逐批产出避免整文件驻留内存；全局 10 万行上限。"""
    conn = open_connection(ds, db=db_name or None, read_timeout=None)
    total_rows = 0
    truncated = False
    try:
        cur = conn.cursor()
        try:
            yield b'{"sections":['
            first_section = True
            for stmt in _split_sql(sql_text):
                if total_rows >= EXPORT_ROW_CAP:
                    truncated = True
                    break
                cur.execute(stmt)
                if cur.description is None:
                    continue
                columns = [col[0] for col in cur.description]
                prefix = b"," if not first_section else b""
                first_section = False
                yield (
                    prefix
                    + b'{"columns":'
                    + json.dumps(columns, ensure_ascii=False).encode("utf-8")
                    + b',"rows":['
                )
                first_row = True
                for row in cur:
                    if total_rows >= EXPORT_ROW_CAP:
                        truncated = True
                        break
                    piece = json.dumps(
                        [_jsonable(v) for v in row], ensure_ascii=False, default=str
                    ).encode("utf-8")
                    yield (b"," if not first_row else b"") + piece
                    first_row = False
                    total_rows += 1
                yield b"]}"
        finally:
            cur.close()
        yield b"]"
        yield b',"truncated":true}' if truncated else b"}"
    finally:
        conn.close()


def _xlsx_stream(ds: DataSource, db_name: Optional[str], sql_text: str) -> Iterator[bytes]:
    """重放 SQL 导出 Excel（G20，openpyxl write_only；单文件一次性产出）：
    每个查询语句一个 sheet（result1..N）；全局 10 万行上限，超限追加 note sheet。"""
    from openpyxl import Workbook  # 局部导入：非 Excel 导出不加载 openpyxl

    wb = Workbook(write_only=True)
    total_rows = 0
    truncated = False
    conn = open_connection(ds, db=db_name or None, read_timeout=None)
    try:
        cur = conn.cursor()
        try:
            sec = 0
            for stmt in _split_sql(sql_text):
                if total_rows >= EXPORT_ROW_CAP:
                    truncated = True
                    break
                cur.execute(stmt)
                if cur.description is None:
                    continue
                sec += 1
                ws = wb.create_sheet(f"result{sec}")
                ws.append([col[0] for col in cur.description])
                for row in cur:
                    if total_rows >= EXPORT_ROW_CAP:
                        truncated = True
                        break
                    ws.append([_jsonable(v) for v in row])
                    total_rows += 1
        finally:
            cur.close()
        if truncated:
            ws = wb.create_sheet("note")
            ws.append([f"超出导出上限 {EXPORT_ROW_CAP} 行，请改用 C22 文件读取/同步通道"])
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        buf.close()  # 及时释放缓冲
    finally:
        conn.close()
    yield data


_EXPORT_FORMATS = {
    "csv": ("text/csv; charset=utf-8", ".csv", _csv_stream),
    "json": ("application/json; charset=utf-8", ".json", _json_stream),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx", _xlsx_stream),
}


@router.get("/export")
def ide_export(
    historyId: int,
    format: str = "csv",
    user: User = Depends(require_perm("run_instance")),
    db: Session = Depends(get_db),
):
    """按 historyId 重放 SQL 导出（G20）：format=csv（默认，BOM）/json/xlsx；导出动作标记 exported。"""
    media_type, ext, stream_fn = _EXPORT_FORMATS.get(format) or (None, None, None)
    if stream_fn is None:
        raise ApiError(PARAM_INVALID, f"format 仅支持 {'/'.join(_EXPORT_FORMATS)}", status=400)
    hist = db.get(IdeHistory, historyId)
    if hist is None:
        raise ApiError(HISTORY_NOT_FOUND, status=404)
    ds = db.get(DataSource, hist.datasource_id)
    if ds is None:
        raise ApiError(DS_NOT_FOUND, "数据源已删除，无法导出", status=404)
    hist.exported = True
    db.commit()
    logger.info("IDE 导出: history=%d format=%s（操作人 %s）", hist.id, format, user.user_name)
    headers = {"Content-Disposition": f"attachment; filename=ide_export_{hist.id}{ext}"}
    return StreamingResponse(
        stream_fn(ds, hist.db_name, hist.rendered_sql or hist.sql_text or ""),
        media_type=media_type,
        headers=headers,
    )


# ---------- 命名脚本 CRUD（I10 G16：用户隔离 + 用户内名称唯一） ----------


class ScriptBody(BaseModel):
    name: str
    datasourceId: Optional[int] = None
    dbName: Optional[str] = None
    content: str = ""


def _script_payload(row: IdeScript) -> dict:
    return {
        "id": row.id,
        "name": row.name,
        "datasourceId": row.datasource_id,
        "dbName": row.db_name,
        "content": row.content or "",
        "createTime": fmt_dt(row.create_time),
        "updateTime": fmt_dt(row.update_time),
    }


def _get_own_script(db: Session, script_id: int, user: User) -> IdeScript:
    """取当前用户自己的脚本（不存在/非本人 → 404，避免越权探测）。"""
    row = db.get(IdeScript, script_id)
    if row is None or row.user_id != user.id:
        raise ApiError(SCRIPT_NOT_FOUND, status=404)
    return row


@router.get("/scripts")
def list_scripts(
    keyword: Optional[str] = None,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """当前用户脚本清单（keyword 匹配脚本名；按更新时间倒序）。"""
    query = db.query(IdeScript).filter(IdeScript.user_id == user.id)
    if keyword:
        query = query.filter(IdeScript.name.like(f"%{keyword}%"))
    rows = query.order_by(IdeScript.update_time.desc()).all()
    return ok([_script_payload(row) for row in rows])


@router.post("/scripts")
def create_script(
    body: ScriptBody,
    user: User = Depends(require_perm("run_instance")),
    db: Session = Depends(get_db),
):
    """保存脚本为新条目（用户内名称唯一）。"""
    name = (body.name or "").strip()
    if not name:
        raise ApiError(PARAM_INVALID, "脚本名不能为空", status=400)
    if db.query(IdeScript).filter(IdeScript.user_id == user.id, IdeScript.name == name).first():
        raise ApiError(PARAM_INVALID, f"脚本名已存在: {name}", status=400)
    row = IdeScript(
        user_id=user.id,
        name=name[:128],
        datasource_id=body.datasourceId,
        db_name=body.dbName,
        content=body.content or "",
    )
    db.add(row)
    db.commit()
    logger.info("IDE 脚本保存: id=%d %s（操作人 %s）", row.id, row.name, user.user_name)
    return ok(_script_payload(row))


@router.put("/scripts/{script_id}")
def update_script(
    script_id: int,
    body: ScriptBody,
    user: User = Depends(require_perm("run_instance")),
    db: Session = Depends(get_db),
):
    """更新脚本（改名时同样校验用户内唯一）。"""
    row = _get_own_script(db, script_id, user)
    name = (body.name or "").strip()
    if not name:
        raise ApiError(PARAM_INVALID, "脚本名不能为空", status=400)
    dup = db.query(IdeScript).filter(IdeScript.user_id == user.id, IdeScript.name == name)
    if script_id:
        dup = dup.filter(IdeScript.id != script_id)
    if dup.first() is not None:
        raise ApiError(PARAM_INVALID, f"脚本名已存在: {name}", status=400)
    row.name = name[:128]
    row.datasource_id = body.datasourceId
    row.db_name = body.dbName
    row.content = body.content or ""
    db.commit()
    logger.info("IDE 脚本更新: id=%d %s（操作人 %s）", row.id, row.name, user.user_name)
    return ok(_script_payload(row))


@router.delete("/scripts/{script_id}")
def delete_script(
    script_id: int,
    user: User = Depends(require_perm("run_instance")),
    db: Session = Depends(get_db),
):
    """删除脚本（仅本人）。"""
    row = _get_own_script(db, script_id, user)
    db.delete(row)
    db.commit()
    logger.info("IDE 脚本删除: id=%d %s（操作人 %s）", row.id, row.name, user.user_name)
    return ok(True)
